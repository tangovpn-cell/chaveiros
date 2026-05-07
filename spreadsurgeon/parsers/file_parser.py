"""Parser universal: CSV, TSV, TXT, XLSX, XLSM, XLS e dados colados."""

import csv
import io
import os
import re
from dataclasses import dataclass, field
from typing import Optional

from spreadsurgeon.utils.text_utils import safe_str, is_scientific_notation

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[dict]
    source_type: str        # csv | tsv | xlsx | xlsm | xls | pasted | txt
    sheet_name: str = ""
    encoding_detected: str = "utf-8"
    warnings: list[str] = field(default_factory=list)
    raw_preview: list[list] = field(default_factory=list)


@dataclass
class ParseResult:
    tables: list[ParsedTable] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    source_file: str = ""
    total_rows: int = 0


def detect_separator(sample: str) -> str:
    counts = {";": sample.count(";"), ",": sample.count(","),
              "\t": sample.count("\t"), "|": sample.count("|")}
    return max(counts, key=counts.get)


def detect_encoding(filepath: str) -> str:
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                f.read(4096)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return "latin-1"


def _sanitize_headers(headers: list[str]) -> list[str]:
    seen = {}
    result = []
    for h in headers:
        h = safe_str(h).strip().lower().replace(" ", "_")
        if not h:
            h = "coluna_vazia"
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 0
        result.append(h)
    return result


def _fix_scientific(value: str) -> str:
    if is_scientific_notation(value):
        try:
            return str(int(float(value)))
        except (ValueError, OverflowError):
            return value
    return value


def _row_to_dict(headers: list[str], row: list) -> dict:
    d = {}
    for i, h in enumerate(headers):
        val = row[i] if i < len(row) else ""
        val = safe_str(val)
        val = _fix_scientific(val)
        d[h] = val
    return d


def parse_csv(content: str, separator: str = None,
              encoding: str = "utf-8") -> ParsedTable:
    lines = content.splitlines()
    if not lines:
        return ParsedTable([], [], "csv", warnings=["Arquivo CSV vazio"])
    sample = "\n".join(lines[:5])
    sep = separator or detect_separator(sample)
    reader = csv.reader(io.StringIO(content), delimiter=sep)
    raw = list(reader)
    if not raw:
        return ParsedTable([], [], "csv", warnings=["CSV sem linhas"])
    headers = _sanitize_headers(raw[0])
    rows = [_row_to_dict(headers, r) for r in raw[1:] if any(c.strip() for c in r)]
    warnings = []
    if sep == ";":
        pass  # Bling padrão
    return ParsedTable(headers=headers, rows=rows, source_type="csv",
                       encoding_detected=encoding,
                       raw_preview=raw[:7], warnings=warnings)


def parse_file(filepath: str) -> ParseResult:
    result = ParseResult(source_file=filepath)
    ext = os.path.splitext(filepath)[1].lower()

    if ext in (".csv", ".txt", ".tsv"):
        encoding = detect_encoding(filepath)
        with open(filepath, "r", encoding=encoding, errors="replace") as f:
            content = f.read()
        sep = "\t" if ext == ".tsv" else None
        table = parse_csv(content, separator=sep, encoding=encoding)
        table.source_type = ext.lstrip(".")
        result.tables.append(table)
        result.total_rows = len(table.rows)

    elif ext in (".xlsx", ".xlsm"):
        if not HAS_OPENPYXL:
            result.errors.append("openpyxl não instalado. Execute: pip install openpyxl")
            return result
        result = _parse_openpyxl(filepath, result, ext.lstrip("."))

    elif ext == ".xls":
        if not HAS_XLRD:
            result.errors.append("xlrd não instalado. Execute: pip install xlrd")
            return result
        result = _parse_xlrd(filepath, result)

    else:
        result.errors.append(f"Formato não suportado: '{ext}'")

    return result


def _parse_openpyxl(filepath: str, result: ParseResult, src_type: str) -> ParseResult:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        raw = [[safe_str(cell.value) for cell in row] for row in ws.iter_rows()]
        raw = [r for r in raw if any(c.strip() for c in r)]
        if not raw:
            continue
        headers = _sanitize_headers(raw[0])
        rows = [_row_to_dict(headers, r) for r in raw[1:]]
        warnings = []
        for i, r in enumerate(raw[1:], 2):
            for j, v in enumerate(r):
                if is_scientific_notation(v):
                    warnings.append(
                        f"Linha {i}, col {headers[j] if j < len(headers) else j}: "
                        f"notação científica detectada → '{v}' → corrigido para '{_fix_scientific(v)}'"
                    )
        table = ParsedTable(headers=headers, rows=rows, source_type=src_type,
                            sheet_name=sheet_name, raw_preview=raw[:7],
                            warnings=warnings)
        result.tables.append(table)
    result.total_rows = sum(len(t.rows) for t in result.tables)
    return result


def _parse_xlrd(filepath: str, result: ParseResult) -> ParseResult:
    wb = xlrd.open_workbook(filepath)
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        raw = []
        for rx in range(ws.nrows):
            row = [safe_str(ws.cell_value(rx, cx)) for cx in range(ws.ncols)]
            raw.append(row)
        raw = [r for r in raw if any(c.strip() for c in r)]
        if not raw:
            continue
        headers = _sanitize_headers(raw[0])
        rows = [_row_to_dict(headers, r) for r in raw[1:]]
        table = ParsedTable(headers=headers, rows=rows, source_type="xls",
                            sheet_name=sheet_name, raw_preview=raw[:7])
        result.tables.append(table)
    result.total_rows = sum(len(t.rows) for t in result.tables)
    return result


def parse_pasted(text: str) -> ParseResult:
    """Interpreta dados colados diretamente (Excel, tabela, CSV inline)."""
    result = ParseResult(source_file="<colado>")
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        result.errors.append("Nenhum dado detectado no texto colado")
        return result
    sample = "\n".join(lines[:5])
    sep = detect_separator(sample)
    table = parse_csv(text, separator=sep, encoding="utf-8")
    table.source_type = "pasted"
    result.tables.append(table)
    result.total_rows = len(table.rows)
    return result
